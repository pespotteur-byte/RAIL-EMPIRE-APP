#!/usr/bin/env python3
"""Post-process catalog: recompute price/tonnage and regenerate Excel."""
import json, re
from pathlib import Path
from collections import Counter

try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
except Exception:
    Workbook = PatternFill = None

REPO = Path('/home/ubuntu/repos/RAIL-EMPIRE-APP')
CATALOG_PATH = REPO / 'js/catalog-data.js'
EXCEL_PATH = REPO / 'catalog_mlg.xlsx'

def read_catalog():
    text = CATALOG_PATH.read_text(encoding='utf-8')
    ct_start = text.find('export const CATALOG_CARGO_TYPES')
    ct_end = text.find('];', text.find('[', ct_start)) + 1
    cat_start = text.find('export const CATALOG = [')
    cat_br = text.find('[', cat_start)
    cat_end = text.rfind('];') + 1
    prefix = text[:ct_start]
    suffix = text[cat_end:]
    cargo = json.loads(text[text.find('[', ct_start):ct_end])
    catalog = json.loads(text[cat_br:cat_end])
    return cargo, catalog, prefix, suffix

def write_catalog(cargo, catalog, prefix, suffix):
    with open(CATALOG_PATH, 'w', encoding='utf-8') as f:
        f.write(prefix)
        f.write('export const CATALOG_CARGO_TYPES = ')
        f.write(json.dumps(cargo, indent=2, ensure_ascii=False))
        f.write(';\n\n')
        f.write('export const CATALOG = ')
        f.write(json.dumps(catalog, indent=2, ensure_ascii=False))
        f.write(';')
        f.write(suffix)

def recompute_entry(e):
    cat = e['category']
    if cat in ('locomotive', 'automotrice'):
        e['purchasePrice'] = int(e.get('power', 0) * 1000)
        e['tonnage'] = e.get('mass', 0)
    elif cat == 'wagon':
        fc = e.get('freightCapacity', 0) or 0
        e['purchasePrice'] = int(fc * 100)
        e['tonnage'] = fc
    elif cat == 'voiture':
        pc = e.get('passengerCapacity', 0) or 0
        e['purchasePrice'] = int(pc * 100)
        e['tonnage'] = pc

def status_color(e):
    cat = e['category']
    if cat in ('locomotive', 'automotrice'):
        has_all = e['maxSpeed'] and e['mass'] and e['power']
        has_some = e['maxSpeed'] or e['mass'] or e['power']
    elif cat == 'wagon':
        has_all = e['maxSpeed'] and e['mass'] and e['freightCapacity']
        has_some = e['maxSpeed'] or e['mass'] or e['freightCapacity']
    elif cat == 'voiture':
        has_all = e['maxSpeed'] and e['mass'] and e['passengerCapacity']
        has_some = e['maxSpeed'] or e['mass'] or e['passengerCapacity']
    else:
        has_all = e['maxSpeed'] and e['mass']
        has_some = e['maxSpeed'] or e['mass']
    if has_all:
        return 'Vert'
    if has_some:
        return 'Jaune/Orange'
    return 'Rouge'

def country_from_image(image):
    parts = image.split('/')
    if len(parts) > 2 and parts[1] == 'catalog':
        return parts[2]
    return ''

def main():
    cargo, catalog, prefix, suffix = read_catalog()
    print(f'Loaded {len(catalog)} entries')
    rows = []
    for e in catalog:
        recompute_entry(e)
        stat = status_color(e)
        rows.append({
            'ID': e['id'],
            'Pays': country_from_image(e.get('imageData','')),
            'Nom': e['name'],
            'Image': Path(e.get('imageData','')).name,
            'Catégorie': e['category'],
            'Traction': e['traction'],
            'Vitesse': e['maxSpeed'],
            'Masse': e['mass'],
            'Puissance': e['power'],
            'Longueur': e['length'],
            'Places': e['passengerCapacity'],
            'Tonnage': e['freightCapacity'],
            'Prix': e['purchasePrice'],
            'Source': e.get('_source', 'MLG Traffic'),
            'Statut': stat,
            'seriesName': e.get('seriesName', ''),
            'cargoTypes': ', '.join(e.get('cargoTypes', [])),
        })
    write_catalog(cargo, catalog, prefix, suffix)
    print('Catalog written')

    if Workbook is None:
        print('openpyxl not available, skipping Excel')
        return
    wb = Workbook()
    ws = wb.active
    ws.title = 'Catalogue MLG'
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        color_map = {'Vert': 'C6EFCE', 'Jaune/Orange': 'FFEB9C', 'Rouge': 'FFC7CE'}
        for r in rows:
            ws.append([r[h] for h in headers])
            fill = PatternFill(start_color=color_map.get(r['Statut'], 'FFFFFF'),
                               end_color=color_map.get(r['Statut'], 'FFFFFF'),
                               fill_type='solid')
            ws.cell(row=ws.max_row, column=headers.index('Statut') + 1).fill = fill
    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(EXCEL_PATH)
    stats = Counter(r['Statut'] for r in rows)
    print('Excel written', EXCEL_PATH, stats)

if __name__ == '__main__':
    main()
