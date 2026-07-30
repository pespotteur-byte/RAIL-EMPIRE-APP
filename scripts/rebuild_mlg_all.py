#!/usr/bin/env python3
# Reconstruction complète du catalogue MLG Traffic (tous pays) + recherche web
# Usage: python3 scripts/rebuild_mlg_all.py
# Génère js/catalog-data.js et catalog_mlg.xlsx

import json, re, os, sys, math, time, urllib.request, urllib.error, urllib.parse, html, threading
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed
from PIL import Image
from io import BytesIO
from bs4 import BeautifulSoup

REPO = Path('/home/ubuntu/repos/RAIL-EMPIRE-APP')
IMG_ROOT = REPO / 'img/catalog'
CATALOG_PATH = REPO / 'js/catalog-data.js'
EXCEL_PATH = REPO / 'catalog_mlg.xlsx'
CACHE_DIR = REPO / '.cache_rebuild'
MLG_BASE = 'http://www.mlgtraffic.net'
# Wikipedia/DB Cargo spec enhancement is disabled for this offline pass due to rate limiting;
# it will be re-enabled/re-run in a separate pass once network conditions allow.
ENABLE_WIKIPEDIA = False
ENABLE_DBCARGO = False

# Country hints for Wikipedia queries based on collection prefix / page text
COUNTRY_HINTS = {
    'SNCF': 'fr', 'France': 'fr', 'RATP': 'fr', 'RegioRail': 'fr', 'VFLI': 'fr',
    'DB': 'de', 'DR': 'de', 'DPost': 'de', 'D_': 'de', 'DSB': 'da', 'DK': 'da',
    'FS': 'it', 'Mercitalia': 'it', 'FNM': 'it', 'SNFT': 'it', 'Trenord': 'it', 'FER': 'it', 'FBP': 'it', 'FP': 'it', 'FSF': 'it', 'GTT': 'it', 'ST': 'it', 'FSE': 'it', 'I_P': 'it', 'I_WP': 'it',
    'CFF': 'fr', 'BLS': 'de', 'SOB': 'de', 'BT': 'de', 'RM': 'de', 'TRN': 'de', 'SZU': 'de', 'GFM': 'de', 'Travys': 'fr', 'MThB': 'de', 'Thurbo': 'de',
    'OBB': 'de', 'A_': 'de', 'PKP': 'pl', 'Rus': 'ru', 'Elipsos': 'es', 'Lyria': 'fr', 'CNL': 'de',
    'SNCB': 'fr', 'NS': 'nl', 'CFL': 'fr', 'NL_': 'nl', 'EETC': 'nl', 'DBNL': 'nl', 'Captrain': 'fr', 'CFLCargo': 'fr', 'Intl': 'en',
    'Constructeurs': 'en', 'Ermewa': 'en', 'MRCE': 'de', 'Nacco': 'en', 'VTG': 'de', 'Loueurs': 'en',
    'Urbains': 'fr',
}

CACHE_DIR.mkdir(exist_ok=True)
fetch_lock = threading.Lock()


def log(msg):
    print(msg, flush=True)

def slugify(s):
    s = re.sub(r'[^\w\s]', ' ', s)
    return re.sub(r'\s+', ' ', s).strip().lower()

def norm_digits(s):
    return set(re.findall(r'\d+', str(s)))

def read_text(path):
    with open(path, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read()

def write_text(path, text):
    with open(path, 'w', encoding='utf-8') as f:
        f.write(text)

DOMAIN_LAST_FETCH = {}
DOMAIN_MIN_DELAY = {
    'wikipedia.org': 0.8,
    'wikimedia.org': 0.8,
    'trains-europe.fr': 0.2,
    'gueterwagenkatalog.dbcargo.com': 0.2,
}
DOMAIN_LOCKS = {}

def _throttle(url):
    domain = urllib.parse.urlparse(url).netloc.lower()
    # strip leading www.
    domain = re.sub(r'^www\.', '', domain)
    delay = DOMAIN_MIN_DELAY.get(domain, 0.05)
    with fetch_lock:
        lock = DOMAIN_LOCKS.get(domain)
        if lock is None:
            lock = threading.Lock()
            DOMAIN_LOCKS[domain] = lock
    with lock:
        now = time.time()
        last = DOMAIN_LAST_FETCH.get(domain, 0)
        wait = max(0, last + delay - now)
        if wait:
            time.sleep(wait)
        DOMAIN_LAST_FETCH[domain] = time.time()

def fetch(url, timeout=30):
    """Fetch with caching (thread-safe)."""
    safe = re.sub(r'[^\w\-_.]', '_', url)
    cache = CACHE_DIR / safe
    with fetch_lock:
        if cache.exists() and cache.stat().st_size > 0:
            return read_text(cache)
    try:
        _throttle(url)
        req = urllib.request.Request(url, headers={'User-Agent':'DevinBot/1.0 (research; devin@example.com)'})
        data = urllib.request.urlopen(req, timeout=timeout).read()
        text = data.decode('utf-8', errors='ignore')
        with fetch_lock:
            # write only if still missing (another thread may have written)
            if not cache.exists():
                write_text(cache, text)
        return text
    except Exception as e:
        log(f'  fetch fail {url}: {e}')
        return None

def fetch_bytes(url, timeout=20):
    try:
        _throttle(url)
        req = urllib.request.Request(url, headers={'User-Agent':'DevinBot/1.0 (research; devin@example.com)'})
        return urllib.request.urlopen(req, timeout=timeout).read()
    except Exception:
        return None

def load_catalog_header():
    """Load CATALOG_CARGO_TYPES header from existing file or default."""
    if CATALOG_PATH.exists():
        text = read_text(CATALOG_PATH)
        m = re.search(r'export const CATALOG_CARGO_TYPES = \[.*?\];', text, re.S)
        if m:
            return m.group(0)
    # fallback minimal cargo types
    return '''// AUTO-GENERATED — Catalogue complet MLG Traffic
export const CATALOG_CARGO_TYPES = [
  {"category":"marchandises","type":"general","name":"Marchandises générales","unit":"t","pricePerUnit":20,"hazard":false}
];
'''

# ---------------------------------------------------------------------------
# Trains d'Europe index & spec resolver
# ---------------------------------------------------------------------------

te_indexes = {}

def build_te_indexes():
    """Fetch all Trains d'Europe indexes and build a slug map."""
    indexes = {
        'locomotives': 'http://trains-europe.fr/sncf/locomotives/index.htm',
        'automoteurs': 'http://trains-europe.fr/sncf/automoteurs/index.htm',
        'tgv': 'http://trains-europe.fr/sncf/tgv/index.htm',
        'turbo': 'http://trains-europe.fr/sncf/turbo/index.htm',
        'voitures': 'http://trains-europe.fr/sncf/voitures/index.htm',
    }
    for name, url in indexes.items():
        text = fetch(url)
        if not text:
            continue
        soup = BeautifulSoup(text, 'html.parser')
        entries = []
        for a in soup.find_all('a', href=True):
            href = a['href'].strip()
            txt = a.get_text(' ', strip=True)
            if not (href.endswith('.htm') or href.endswith('.html')):
                continue
            if href.startswith('..') or href.startswith('http'):
                continue
            if 'index' in href.lower():
                continue
            # resolve relative to base url
            base = url.rsplit('/', 1)[0]
            full = f'{base}/{href}'
            entries.append({'title': txt, 'slug': full, 'href': href})
        te_indexes[name] = entries
        log(f'TE index {name}: {len(entries)} pages')

def te_find_slug(name, category):
    """Find best matching Trains d'Europe slug for a given name/category."""
    if not name:
        return None
    norm = slugify(name)
    tokens = set(norm.split())
    digits = norm_digits(name)
    candidates = []
    if category in ('locomotive', 'automotrice', 'tgv'):
        for idx_name in (('tgv' if 'tgv' in norm else None),
                         ('turbo' if 'turbo' in norm or 'rtg' in norm or 'etg' in norm else None),
                         ('automoteurs' if category == 'automotrice' else None),
                         ('locomotives' if category in ('locomotive','automotrice') else None)):
            if not idx_name:
                continue
            for e in te_indexes.get(idx_name, []):
                title_norm = slugify(e['title'])
                title_digits = norm_digits(e['title'])
                score = 0
                # exact class token match
                if norm and (norm in title_norm or title_norm in norm):
                    score += 50
                # digit overlap
                if digits and title_digits and digits & title_digits:
                    score += 20
                # keyword matches
                for tok in tokens:
                    if tok in title_norm:
                        score += 5
                # prefer TGV index for TGV
                if idx_name == 'tgv' and 'tgv' in norm:
                    score += 30
                if score:
                    candidates.append((score, e))
    else:
        for e in te_indexes.get('voitures', []):
            title_norm = slugify(e['title'])
            title_digits = norm_digits(e['title'])
            score = 0
            if norm in title_norm or title_norm in norm:
                score += 50
            if digits and title_digits and digits & title_digits:
                score += 20
            for tok in tokens:
                if tok in title_norm:
                    score += 5
            if score:
                candidates.append((score, e))
    if not candidates:
        return None
    candidates.sort(key=lambda x: x[0], reverse=True)
    return candidates[0][1]['slug']

te_spec_cache = {}
te_cache_lock = threading.Lock()

def parse_te_page(url):
    """Parse a Trains d'Europe page into a list of fiche dicts."""
    with te_cache_lock:
        if url in te_spec_cache:
            return te_spec_cache[url]
    text = fetch(url)
    if not text:
        return []
    soup = BeautifulSoup(text, 'html.parser')
    fiches = []
    for fiche in soup.find_all(['div', 'td'], class_=lambda c: c and 'fiche' in c):
        data = {}
        for d in fiche.find_all('div', style=lambda s: s and 'display:inline' in s):
            strong = d.find('div', style=lambda s: s and 'font-weight:bold' in s)
            if not strong:
                continue
            label = strong.get_text(strip=True).rstrip(':').rstrip()
            value = d.get_text(' ', strip=True).replace(strong.get_text(strip=True), '', 1).strip()
            data[label] = value
        if data:
            fiches.append(data)
    with te_cache_lock:
        te_spec_cache[url] = fiches
    return fiches

def norm_num_str(s):
    """Collapse all visual spacing characters between digits."""
    s2 = str(s)
    # collapse unicode non-breaking spaces / thin spaces between digits
    for ch in '\u00a0\u2007\u202f\u2009\u2006\u2002\u2003\u2008\u205f':
        s2 = s2.replace(ch, ' ')
    s2 = re.sub(r'(?<=\d)\s+(?=\d)', '', s2)
    return s2

def first_num(s):
    if not s:
        return None
    s2 = norm_num_str(s)
    m = re.search(r'[\d\s]+(?:[,.][\d\s]+)?', s2)
    if not m:
        return None
    v = m.group(0).replace(' ', '').replace(',', '.')
    try:
        return float(v)
    except Exception:
        return None

def all_nums(s):
    if not s:
        return []
    s2 = norm_num_str(s)
    vals = []
    for m in re.finditer(r'\d[\d\s,\.]*\d|\d', s2):
        v = m.group(0).replace(' ', '').replace(',', '.')
        try:
            vals.append(float(v))
        except Exception:
            pass
    return vals

def select_multi_value(text, car_count, unit_patterns=None):
    """For specs like 'si rame à 2 caisses : 94 t.', return value for car_count."""
    if not text or not car_count:
        return None
    text = re.sub(r'\s+', ' ', text).strip()
    pattern = r'(?:si\s+)?(?:rame\s+à\s+)?(\d+)\s*(?:caisses?|bodies?|voitures?)\s*[:：]\s*([^;\n]*?)(?=\s*(?:si\s+rame|(?:\d+)\s*(?:caisses?|bodies?|voitures?)\s*[:：]|$))'
    matches = re.findall(pattern, text, re.I)
    if matches:
        for count, val in matches:
            if int(count) == car_count:
                return val.strip()
    return None

def pick_fiche(fiches, name, niv2):
    """Pick the best fiche matching a subseries."""
    if not fiches:
        return None
    if len(fiches) == 1:
        return fiches[0]
    # try to match by series indicators in name/niv2
    series_indicators = []
    for m in re.finditer(r'(1[èe]re?|premi[èe]re|2[èe]me|deuxi[èe]me|3[èe]me|troisi[èe]me|4[èe]me)', str(name)+' '+str(niv2), re.I):
        series_indicators.append(m.group(1).lower()[:3])
    if not series_indicators:
        return fiches[0]
    # fiche contains Numéros or text
    for i, f in enumerate(fiches):
        nums = str(f.get('Numéros', '')) + ' ' + str(f.get('Nombre d\'exemplaires', ''))
        if series_indicators[0] in nums.lower():
            return f
    return fiches[0]

def parse_spec_value(label, value, car_count=None):
    """Extract numeric / text fields from a fiche."""
    if not value:
        return {}
    res = {}
    vstr = str(value)
    if label in ('Puissance',):
        nums = all_nums(vstr)
        if nums:
            res['power_kw'] = max(nums) if 'sous' in vstr.lower() or any(x in vstr for x in ['ou','et',';']) else nums[0]
    elif label in ('Poids', 'Masse', 'Masse en service', 'Masse totale'):
        sel = select_multi_value(vstr, car_count)
        nums = all_nums(sel or vstr)
        if nums:
            res['mass_t'] = nums[0]
    elif label in ('Vitesse maximale', 'Vitesse'):
        nums = all_nums(vstr)
        if nums:
            res['max_speed_kmh'] = max(nums)
    elif label in ('Longueur', 'Longueur hors-tout', 'Longueur totale'):
        sel = select_multi_value(vstr, car_count)
        nums = all_nums(sel or vstr)
        if nums:
            res['length_m'] = nums[0]
    elif label in ('Places de première classe',):
        nums = all_nums(vstr)
        if nums:
            res['first_class'] = int(nums[0])
    elif label in ('Places de seconde classe',):
        nums = all_nums(vstr)
        if nums:
            res['second_class'] = int(nums[0])
    elif label in ('Places de troisième classe','Places de 3e classe'):
        nums = all_nums(vstr)
        if nums:
            res['third_class'] = int(nums[0])
    elif label in ('Places', 'Capacité'):
        nums = all_nums(vstr)
        if nums:
            res['capacity'] = int(nums[0])
    return res

def merge_specs(fiche, name, niv2, car_count=None):
    spec = {}
    if fiche:
        for label, val in fiche.items():
            spec.update(parse_spec_value(label, val, car_count))
    return spec

def resolve_specs_te(name, niv1, niv2, category, traction, car_count=None):
    """Try to resolve specs from Trains d'Europe."""
    # Use Niv1 title for TGV, otherwise name
    search_name = name if category != 'tgv' else (niv1 or name)
    slug = te_find_slug(search_name, category)
    if not slug:
        return {}
    fiches = parse_te_page(slug)
    if not fiches:
        return {}
    f = pick_fiche(fiches, name, niv2)
    spec = merge_specs(f, name, niv2, car_count)
    if not spec:
        return {}
    spec['_source_te'] = slug
    return spec

# ---------------------------------------------------------------------------
# Wikipedia fallback (for non-wagons, multiple languages)
# ---------------------------------------------------------------------------

WIKI_LABELS = {
    'fr': {
        'power': {'puissance'},
        'mass': {'masse en service', 'masse', 'poids'},
        'speed': {'vitesse maximale', 'vitesse de ligne'},
        'length': {'longueur', 'longueur hors-tout'},
        'passenger': {'places', 'capacité'},
        'first': {"places de première classe", "1re classe"},
        'second': {"places de seconde classe", "2e classe"},
        'third': {"places de troisième classe", "3e classe"},
    },
    'en': {
        'power': {'power output', 'power rating', 'continuous power', 'power'},
        'mass': {'service weight', 'weight on drivers', 'loco weight', 'weight', 'mass'},
        'speed': {'maximum speed', 'max speed', 'top speed', 'service speed'},
        'length': {'length', 'overall length', 'over buffers'},
        'passenger': {'seating', 'seating capacity', 'passenger capacity', 'capacity'},
        'first': {'1st class', 'first class'},
        'second': {'2nd class', 'second class'},
        'third': {'3rd class', 'third class'},
    },
    'de': {
        'power': {'leistung', 'nennleistung', 'stundenleistung'},
        'mass': {'dienstmasse', 'masse', 'gewicht', 'reibungsmasse'},
        'speed': {'höchstgeschwindigkeit', 'geschwindigkeit'},
        'length': {'länge', 'länge über puffer'},
        'passenger': {'sitze', 'sitzplätze', 'plätze'},
        'first': {'1. klasse', 'erste klasse'},
        'second': {'2. klasse', 'zweite klasse'},
        'third': {'3. klasse', 'dritte klasse'},
    },
    'it': {
        'power': {'potenza', 'potenza di trazione'},
        'mass': {'massa', 'massa in servizio', 'peso'},
        'speed': {'velocità massima', 'velocità'},
        'length': {'lunghezza', 'lunghezza totale'},
        'passenger': {'posti a sedere', 'posti', 'capacità'},
        'first': {'prima classe', '1ª classe'},
        'second': {'seconda classe', '2ª classe'},
        'third': {'terza classe', '3ª classe'},
    },
    'es': {
        'power': {'potencia', 'potencia máxima'},
        'mass': {'masa', 'peso'},
        'speed': {'velocidad máxima', 'velocidad'},
        'length': {'longitud', 'longitud total'},
        'passenger': {'capacidad', 'plazas'},
        'first': {'primera clase', '1ª clase'},
        'second': {'segunda clase', '2ª clase'},
        'third': {'tercera clase', '3ª clase'},
    },
}

def _label_match(kl, labels):
    for lab in labels:
        if lab in kl:
            return True
    return False

def _extract_passenger_counts(v, labels):
    """Try to extract 1st/2nd/3rd class seat counts from a cell value."""
    out = {}
    for cls in ('first', 'second', 'third'):
        cls_labels = labels.get(cls, set())
        if not cls_labels:
            continue
        # look for pattern "<label> : 42"
        for lab in cls_labels:
            m = re.search(rf'(?:{re.escape(lab)})\s*[:：]\s*([\d\s,\.]+)', v, re.I)
            if m:
                nums = all_nums(m.group(1))
                if nums:
                    out[cls] = int(nums[0])
                    break
    return out

def parse_wikipedia_html(text, lang):
    labels = WIKI_LABELS.get(lang, WIKI_LABELS['en'])
    soup = BeautifulSoup(text, 'html.parser')
    spec = {}
    # Infobox tables
    for table in soup.find_all('table'):
        for tr in table.find_all('tr'):
            cells = [c.get_text(' ', strip=True) for c in tr.find_all(['th','td'])]
            if len(cells) < 2:
                continue
            k = cells[0].lower()
            v = ' '.join(cells[1:])
            if _label_match(k, labels['power']):
                nums = all_nums(v)
                if nums:
                    spec['power_kw'] = max(nums)
            if _label_match(k, labels['mass']):
                nums = all_nums(v)
                if nums:
                    spec['mass_t'] = nums[0]
            if _label_match(k, labels['speed']):
                nums = all_nums(v)
                if nums:
                    spec['max_speed_kmh'] = max(nums)
            if _label_match(k, labels['length']):
                nums = all_nums(v)
                if nums:
                    spec['length_m'] = nums[0]
            if _label_match(k, labels['passenger']):
                counts = _extract_passenger_counts(v, labels)
                for cls, n in counts.items():
                    spec[f'{cls}_class'] = n
                nums = all_nums(v)
                if nums and 'capacity' not in spec:
                    spec['capacity'] = int(nums[0])
    # Fallback body regex
    body = soup.get_text(' ', strip=True)
    if 'power_kw' not in spec:
        for pat in [
            r'(?:power output|power)[^.\n]{0,80}?(\d[\d\s,\.]+)\s*(?:kW|kilowatt)',
            r'(\d[\d\s,\.]+)\s*kW',
        ]:
            m = re.search(pat, body, re.I)
            if m:
                nums = all_nums(m.group(1))
                if nums:
                    spec['power_kw'] = max(nums)
                    break
    if 'mass_t' not in spec:
        for pat in [
            r'(?:service weight|weight|mass)[^.\n]{0,80}?(\d[\d\s,\.]+)\s*(?:t|tonnes?)',
            r'(\d[\d\s,\.]+)\s*(?:t|tonnes?)',
        ]:
            m = re.search(pat, body, re.I)
            if m:
                nums = all_nums(m.group(1))
                if nums:
                    spec['mass_t'] = nums[0]
                    break
    if 'max_speed_kmh' not in spec:
        for pat in [
            r'(?:maximum speed|top speed|max speed|service speed)[^.\n]{0,80}?(\d[\d\s,\.]+)\s*(?:km/h|km\s*h)',
            r'(\d[\d\s,\.]+)\s*(?:km/h|km\s*h)',
        ]:
            m = re.search(pat, body, re.I)
            if m:
                nums = all_nums(m.group(1))
                if nums:
                    spec['max_speed_kmh'] = max(nums)
                    break
    return spec

def _wikipedia_search_title(name, lang):
    """Use Wikipedia opensearch API to find the article title for a model."""
    if not name:
        return None
    q = f'{name} train'
    url = f'https://{lang}.wikipedia.org/w/api.php?action=opensearch&search={urllib.parse.quote(q)}&limit=3&format=json'
    text = fetch(url, timeout=20)
    if not text:
        return None
    try:
        data = json.loads(text)
        # data = [query, [titles], [descs], [urls]]
        titles = data[1] if len(data) > 1 else []
        for t in titles:
            t = t.strip()
            if t:
                return t.replace(' ', '_')
    except Exception:
        pass
    return None

def resolve_specs_wikipedia(name, category, langs=None):
    """Try to find specs on Wikipedia in one or more languages."""
    if category == 'wagon':
        return {}
    if not name:
        return {}
    # Skip names that are too short / generic and unlikely to match (e.g. single-letter coaches)
    name_clean = re.sub(r'[^\w\s]', ' ', name).strip()
    if len(name_clean) < 3:
        return {}
    if category == 'voiture' and len(name_clean) < 6 and not re.search(r'\d', name_clean):
        return {}
    if langs is None:
        langs = ['en', 'fr']
    for lang in langs:
        title = _wikipedia_search_title(name, lang)
        if not title:
            # Fallback to direct title guess
            title = name.replace(' ', '_')
        url = f'https://{lang}.wikipedia.org/wiki/{urllib.parse.quote(title)}'
        cache_key = re.sub(r'[^\w]', '_', f'{lang}_{title}')
        cache = CACHE_DIR / f'wiki_{cache_key}'
        if cache.exists() and cache.stat().st_size > 0:
            text = read_text(cache)
        else:
            data = fetch_bytes(url, timeout=20)
            if not data:
                continue
            text = data.decode('utf-8', errors='ignore')
            write_text(cache, text)
        spec = parse_wikipedia_html(text, lang)
        if spec:
            spec['_source_wiki'] = url
            return spec
    return {}

def resolve_specs(name, niv1, niv2, category, traction, car_count, url):
    """Resolve technical specs from best available sources."""
    if not name:
        return {}
    prefix, lang = country_from_url(url)
    specs = {}
    # Trains d'Europe for French rolling stock (locomotives, EMU, coaches)
    if prefix in ('SNCF', 'France', 'RATP'):
        specs = resolve_specs_te(name, niv1, niv2, category, traction, car_count)
    # Wikipedia for locomotives and self-propelled units only (coaches are rarely well covered per variant)
    if ENABLE_WIKIPEDIA and not specs and category in ('locomotive', 'automotrice'):
        langs = []
        if lang and lang != 'en':
            langs.append(lang)
        langs.extend(['en', 'fr'])
        specs = resolve_specs_wikipedia(name, category, langs=langs)
    # For wagons and coaches we do not use Wikipedia; DB Cargo/other sources and manual completion will follow
    return specs or {}

# ---------------------------------------------------------------------------
# MLG parsing
# ---------------------------------------------------------------------------

def _quote_href(href):
    """Quote path but keep query string intact."""
    href = href.split('#')[0]
    if '?' in href:
        path, qs = href.split('?', 1)
        path = urllib.parse.quote(path, safe='/')
        return f'{path}?{qs}'
    return urllib.parse.quote(href, safe='/')

def _extract_page_urls(text):
    """Extract .xml/.htm hrefs from a collection or index page text."""
    urls = set()
    for m in re.finditer(r'href="([^"]+)"', text):
        href = m.group(1)
        if href.lower().startswith(('http:', 'https:', 'mailto:')):
            continue
        if not re.search(r'\.(xml|htm)(\?|$)', href, re.I):
            continue
        # skip CSS and index links
        if href.lower().endswith('.css'):
            continue
        base = href.split('?')[0].lower()
        if base in ('index_fr.html', 'index.html'):
            continue
        urls.add(f'{MLG_BASE}/{_quote_href(href)}')
    return sorted(urls)

def all_page_urls():
    """Return all MLG page URLs from index_fr.html, recursively following collection pages."""
    index = fetch(f'{MLG_BASE}/index_fr.html')
    if not index:
        return []
    # The index page itself contains direct collection/page links
    index_urls = _extract_page_urls(index)
    # Identify collection pages (Coll_*.htm) and direct pages
    collection_urls = [u for u in index_urls if re.search(r'/Coll_[^/]+\.htm', u, re.I)]
    direct_urls = [u for u in index_urls if not re.search(r'/Coll_[^/]+\.htm', u, re.I)]
    all_urls = set(direct_urls)
    for coll_url in collection_urls:
        text = fetch(coll_url)
        if not text:
            continue
        for u in _extract_page_urls(text):
            all_urls.add(u)
    return sorted(all_urls)

def country_from_url(url):
    """Guess country/operator code from an MLG page URL or XML name."""
    name = Path(url).name.split('.')[0]
    # prefer known prefixes before underscore
    for prefix, lang in COUNTRY_HINTS.items():
        if name.startswith(prefix):
            return prefix, lang
    return '', 'en'

def parse_mlg_xml(url, text):
    """Parse an MLG XML file into raw entries. Supports composite Ligne with multiple Image tags."""
    import xml.etree.ElementTree as ET
    root = ET.fromstring(text)
    page_title = ''
    for tag in ['Nompage1', 'Nompage2']:
        el = root.find(tag)
        if el is not None and el.text:
            page_title += ' ' + el.text.strip()
    entries = []
    for niv1 in root.findall('Niv1'):
        t1 = niv1.find('Titre1')
        niv1_title = (t1.text or '').strip() if t1 is not None else ''
        for niv2 in niv1.findall('Niv2'):
            t2 = niv2.find('Titre2')
            niv2_title = (t2.text or '').strip() if t2 is not None else ''
            for ligne in niv2.findall('Ligne'):
                nom_el = ligne.find('Nom')
                notes_el = ligne.find('Notes')
                nom = (nom_el.text or '').strip() if nom_el is not None and nom_el.text else ''
                notes = (notes_el.text or '').strip() if notes_el is not None and notes_el.text else ''
                # Prefer right-facing image list; MLG uses <Image_R>/<Image_L> on some pages
                image_tags = []
                if ligne.find('Image_R') is not None:
                    image_tags = ['Image_R']
                elif ligne.find('Image') is not None:
                    image_tags = ['Image']
                elif ligne.find('Image_L') is not None:
                    image_tags = ['Image_L']
                images = []
                for tag in image_tags:
                    for i in ligne.findall(tag):
                        if i.text and i.text.strip():
                            images.append(i.text.strip())
                if not images:
                    continue
                entries.append({
                    'url': url,
                    'page_title': page_title.strip(),
                    'niv1': niv1_title,
                    'niv2': niv2_title,
                    'nom': nom,
                    'images': images,
                    'typeligne': ligne.get('Typeligne') or '1',
                    'notes': notes,
                })
    return entries

def parse_mlg_htm(url, text):
    """Parse an MLG HTML data page (const dataPage JSON) into raw entries."""
    text = text.replace('\r', '\n')
    m = re.search(r'const\s+dataPage\s*=\s*(\{.*?\}\s*);?\s*<\/script>', text, re.S)
    if not m:
        return []
    raw = m.group(1)
    # remove trailing commas
    raw = re.sub(r',(\s*[}\]])', r'\1', raw)
    try:
        data = json.loads(raw)
    except Exception as e:
        log(f'JSON parse fail {url}: {e}')
        return []
    page_title = ' '.join([str(data.get(k, '')) for k in ['titrePageF','titre2F','titre3F'] if data.get(k)])
    entries = []
    niv1_list = data.get('niv1', [])
    if isinstance(niv1_list, dict):
        niv1_list = [niv1_list]
    for niv1 in niv1_list:
        niv1_title = str(niv1.get('libF', ''))
        niv2_list = niv1.get('niv2', [])
        if isinstance(niv2_list, dict):
            niv2_list = [niv2_list]
        for niv2 in (niv2_list or [niv1]):
            niv2_title = str(niv2.get('libF', '') or niv2.get('txtF', ''))
            mat_list = niv2.get('mat', [])
            if isinstance(mat_list, dict):
                mat_list = [mat_list]
            for mat in mat_list:
                image = str(mat.get('image', ''))
                nom = str(mat.get('nomF', '') or mat.get('txtF', '') or niv1_title)
                notes = str(mat.get('descrF', ''))
                if not image:
                    continue
                entries.append({
                    'url': url,
                    'page_title': page_title.strip(),
                    'niv1': niv1_title,
                    'niv2': niv2_title,
                    'nom': nom,
                    'images': [image],
                    'typeligne': str(mat.get('nbSensMat', '1') or '1'),
                    'notes': notes,
                })
    return entries

# ---------------------------------------------------------------------------
# Image handling
# ---------------------------------------------------------------------------

def resolve_image(base_rel):
    """Download the image and return (side, local_relative_path, width)."""
    base_name = Path(base_rel).name
    sub_dir = Path(base_rel).parent.relative_to('images') if base_rel.startswith('images/') else Path('.')
    local_dir = IMG_ROOT / sub_dir
    local_dir.mkdir(parents=True, exist_ok=True)
    encoded_base = urllib.parse.quote(base_rel, safe='/')

    # If the base itself already ends with a side suffix (e.g. ..._R, ..._Anim_R),
    # treat it as an exact filename base; try .gif and _Anim.gif first.
    has_side_suffix = bool(re.search(r'_(R|L|Anim_R|Anim_L)$', base_name))
    if has_side_suffix:
        priority = [
            ('R', f'{base_name}.gif'),
            ('R', f'{base_name}_Anim.gif'),
        ]
    else:
        # Generic base: right-facing variants first, then non-sided, then left as fallback.
        priority = [
            ('R', f'{base_name}_R.gif'),
            ('R', f'{base_name}_Anim_R.gif'),
            ('', f'{base_name}.gif'),
            ('', f'{base_name}_Anim.gif'),
            ('L', f'{base_name}_L.gif'),
            ('L', f'{base_name}_Anim_L.gif'),
        ]

    for side, filename in priority:
        local = local_dir / filename
        if local.exists() and local.stat().st_size > 0:
            try:
                im = Image.open(local)
                return [(side or 'R', str(local.relative_to(REPO)), im.width)]
            except Exception:
                continue
        suffix = filename[len(base_name):] if filename.startswith(base_name) else filename
        url = f'{MLG_BASE}/{encoded_base}{urllib.parse.quote(suffix, safe="/")}'
        data = fetch_bytes(url, timeout=8)
        if data:
            try:
                im = Image.open(BytesIO(data))
            except Exception:
                continue
            local.write_bytes(data)
            return [(side or 'R', str(local.relative_to(REPO)), im.width)]
    return []

def build_composite(raw, image_results):
    """Build a horizontal composite PNG from a list of resolved (side, path, width) images."""
    parts = []
    widths = []
    for side, img_path, width in image_results:
        p = REPO / img_path
        if not p.exists():
            continue
        try:
            im = Image.open(p).convert('RGBA')
            parts.append(im)
            widths.append(width)
        except Exception:
            continue
    if not parts:
        return None, 0
    total_w = sum(im.width for im in parts)
    max_h = max(im.height for im in parts)
    canvas = Image.new('RGBA', (total_w, max_h), (0, 0, 0, 0))
    x = 0
    for im in parts:
        canvas.paste(im, (x, max_h - im.height), im)
        x += im.width
    # Determine safe output path under the first image's directory/composed
    first_rel = Path(image_results[0][1])
    out_dir = (REPO / first_rel).parent / 'composed'
    out_dir.mkdir(parents=True, exist_ok=True)
    nom = re.sub(r'[^\w\-]', '_', (raw['nom'] or raw['niv1'] or 'comp').strip())[:60]
    notes = re.sub(r'[^\w\-]', '_', (raw['notes'] or '').strip())[:40]
    base = f'{nom}_{notes}' if notes else nom
    base = base.strip('_') or 'comp'
    out_path = out_dir / f'{base}.png'
    counter = 1
    while out_path.exists():
        out_path = out_dir / f'{base}_{counter}.png'
        counter += 1
    try:
        canvas.save(out_path)
    except Exception as e:
        log(f'  composite save error {out_path}: {e}')
        return None, 0
    rel = str(out_path.relative_to(REPO))
    return rel, total_w

# ---------------------------------------------------------------------------
# Cargo type mapping for freight wagons (from rebuild_mlg_catalog.py)
# ---------------------------------------------------------------------------

XML_WAGON_MAP = {
    'Citernes': {'sub': 'citerne', 'cargo': ['oil','chemicals','milk','wine','sulphur','acid','ammonia','lng','lpg']},
    'Plats': {'sub': 'plat', 'cargo': ['steel','steel-coils','steel-sheet','steel-beams','containers','wood','vehicles','sinter','pig-iron','scrap']},
    'CerealB': {'sub': 'tremie', 'cargo': ['grain','flour','sugar','salt','coal','ore','gravel','sand','ballast','limestone','cement','clinker','coke','bauxite','fertilizer','potash','phosphate','kaolin','alumina','sinter','pig-iron','scrap','aggregates']},
    'CerealE': {'sub': 'tremie', 'cargo': ['grain','flour','sugar','salt','coal','ore','gravel','sand','ballast','limestone','cement','clinker','coke','bauxite','fertilizer','potash','phosphate','kaolin','alumina','sinter','pig-iron','scrap','aggregates']},
    'Tremies': {'sub': 'tremie', 'cargo': ['coal','ore','gravel','sand','ballast','grain','flour','sugar','salt','fertilizer','potash','phosphate','kaolin','alumina','sinter','pig-iron','scrap','aggregates','coke','bauxite','limestone','cement','clinker']},
    'Frigo': {'sub': 'frigo', 'cargo': ['refrigerated','meat']},
    'Combis': {'sub': 'combi', 'cargo': ['containers','general','steel','wood']},
    'GEFCO': {'sub': 'vehicles', 'cargo': ['vehicles']},
    'STVA': {'sub': 'vehicles', 'cargo': ['vehicles']},
    'SGW': {'sub': 'plat', 'cargo': ['steel','steel-coils','steel-sheet','containers','wood']},
    'Railtrans': {'sub': 'plat', 'cargo': ['containers','steel','vehicles']},
    'Div': {'sub': 'general', 'cargo': ['general','paper','wood']},
    'EVS': {'sub': 'waste', 'cargo': ['waste']},
    'WEFT': {'sub': 'tremie', 'cargo': ['coal','ore','gravel','sand','ballast','grain','flour','sugar','salt','fertilizer','potash','phosphate','kaolin','alumina','sinter','pig-iron','scrap','aggregates']},
    'WFMVAC': {'sub': 'service', 'cargo': []},
    'WG': {'sub': 'couvert', 'cargo': ['general','paper','wood','sugar','salt','fertilizer','cement','flour','grain']},
    'WHI': {'sub': 'frigo', 'cargo': ['refrigerated','meat','livestock']},
    'WKLRS': {'sub': 'flat', 'cargo': ['steel','steel-coils','steel-sheet','containers','wood']},
}

def guess_wagon_sub(xml_name, image_basename, nom='', series_name=''):
    tail = Path(xml_name).stem.split('_')[-1]
    m = XML_WAGON_MAP.get(tail)
    if m:
        return m['sub']
    low = image_basename.lower()
    name_low = (nom or '').lower()
    series_low = (series_name or '').lower()
    combined = low + ' ' + name_low + ' ' + series_low
    if any(k in combined for k in ['citerne','tank','soufre','gaz','gpl','lng','chemical']):
        return 'citerne'
    if any(k in combined for k in ['cereal','grain','hopper','tremie','trémie']) and 'm3' in combined:
        return 'tremie'
    if any(k in combined for k in ['frigo','frigor','refrig']):
        return 'frigo'
    if 'couvert' in combined or 'covered' in combined:
        return 'couvert'
    if 'combi' in combined or 'contene' in combined:
        return 'combi'
    if any(k in combined for k in ['gefc','stva','auto','vehic']):
        return 'vehicles'
    if any(k in combined for k in ['fourgon','caboose','service','draisine']):
        return 'service'
    if any(k in combined for k in ['waste','evs','dechet']):
        return 'waste'
    if any(k in combined for k in ['plat','open wagon','flat','sgw','klms']):
        return 'plat'
    return 'default'

def get_wagon_cargos(sub):
    for tail, m in XML_WAGON_MAP.items():
        if m['sub'] == sub:
            return list(m['cargo'])
    if sub == 'tremie':
        return ['coal','ore','gravel','sand','ballast','grain','flour','sugar','salt','fertilizer','potash','phosphate','kaolin','alumina','sinter','pig-iron','scrap','aggregates','coke','bauxite','limestone','cement','clinker']
    if sub == 'citerne':
        return ['oil','chemicals','milk','wine','sulphur','acid','ammonia','lng','lpg']
    if sub == 'plat':
        return ['steel','steel-coils','steel-sheet','steel-beams','scrap','containers','vehicles','wood','sinter','pig-iron','general']
    if sub == 'couvert':
        return ['general','paper','wood','sugar','salt','fertilizer','cement','flour','grain']
    if sub == 'frigo':
        return ['refrigerated','meat','livestock']
    if sub == 'combi':
        return ['containers','general','steel','wood']
    if sub == 'vehicles':
        return ['vehicles']
    if sub == 'waste':
        return ['waste']
    return ['general']

# ---------------------------------------------------------------------------
# Categorisation
# ---------------------------------------------------------------------------

def _contains(text, words):
    for w in words:
        if re.search(r'(?<!\w)' + re.escape(w) + r'(?!\w)', text, re.I):
            return True
    return False

def infer_category_traction(page_url, page_title, nom, niv1, niv2, notes):
    path = page_url.split('/')[-1].lower()
    name_l = (nom or '').lower()
    niv1_l = (niv1 or '').lower()
    niv2_l = (niv2 or '').lower()
    notes_l = (notes or '').lower()
    page_l = (page_title or '').lower()
    combined = f'{page_l} {niv1_l} {niv2_l} {name_l} {notes_l} {path}'

    elec_kw = ['électrique', 'electrique', 'panto', 'caténaire', 'sncf_f_le_', '_le_', 'le_',
               'bb ', 'cc ', '1abba1', '1cc1', '2cc2', '2bb2', 'ee ', 're ', 'ae ', 'cable']
    diesel_kw = ['diesel', 'sncf_f_ld_', '_ld_', 'ld_', 'autorail', 'railcar', 'schienenbus',
                 'diesellok']
    steam_kw = ['vapeur', 'steam', 'sncf_f_lv_', '_lv_', 'lv_', 'à vapeur']
    has_elec = _contains(combined, elec_kw) or 'panto' in combined
    has_diesel = _contains(combined, diesel_kw)
    has_steam = _contains(combined, steam_kw)

    if _contains(combined, ['vapeur', 'locomotive à vapeur', 'locomotives à vapeur', 'steam']):
        return 'locomotive', 'vapeur'

    if _contains(combined, ['locomotive', 'locomotives']):
        if has_diesel and not has_elec:
            return 'locomotive', 'diesel'
        if has_steam:
            return 'locomotive', 'vapeur'
        if 'ld_' in path or path.startswith('sncf_f_ld_') or '_ld_' in path:
            return 'locomotive', 'diesel'
        if 'lv_' in path or path.startswith('sncf_f_lv_') or '_lv_' in path:
            return 'locomotive', 'vapeur'
        return 'locomotive', 'electrique'

    emu_kw = ['tgv', 'tvg', 'talis', 'eurostar', 'thalys', 'sncf_f_idf_a', 'ratp',
              'agc', 'regiolis', 'transilien', 'rer', 'métro', 'metro', 'flirt', 'nina',
              'gtw', 'desiro', 'talent', 'a-ter', 'a_ter', 'a ter', 'xr ', 'sncf_f_reg_',
              'sncf_f_priv', 'france_f_trv', 'hexafret']
    if _contains(combined, emu_kw):
        trailer_kw = ['remorque', 'beiwagen', 'sans cabine', 'sans poste de conduite',
                      'ohne führerstand', 'wagen ohne', 'remorques intermédiaires',
                      'remorque intermédiaire']
        pilot_kw = ['pilote', 'pilot', 'voiture pilote', 'cabine de conduite',
                    'poste de conduite', 'driving trailer']
        if _contains(combined, pilot_kw) or ('remorque' in combined and 'pilote' in combined):
            return 'voiture', 'none'
        if _contains(combined, trailer_kw):
            return 'voiture', 'none'
        if has_diesel and not has_elec:
            return 'automotrice', 'diesel'
        return 'automotrice', 'electrique'

    auto_kw = ['automotrice', 'automotrices', 'autorail', 'autorails', 'draisine',
               'train léger', 'schienenbus']
    if _contains(combined, auto_kw):
        trailer_kw = ['remorque', 'beiwagen', 'sans cabine', 'sans poste de conduite',
                      'ohne führerstand', 'wagen ohne']
        pilot_kw = ['pilote', 'pilot', 'voiture pilote', 'poste de conduite',
                    'cabine de conduite', 'driving trailer']
        if _contains(combined, pilot_kw) or ('remorque' in combined and 'pilote' in combined):
            return 'voiture', 'none'
        if _contains(combined, trailer_kw):
            return 'voiture', 'none'
        if has_diesel and not has_elec:
            return 'automotrice', 'diesel'
        if has_steam:
            return 'automotrice', 'vapeur'
        return 'automotrice', 'electrique'

    if _contains(combined, ['voiture', 'voitures', 'coach', 'coaches', 'carriage', 'carriages',
                            'wagen', 'reisezugwagen', 'rijtuig', 'vettura', 'coche', 'coches']):
        if _contains(combined, ['draisine', 'automotrice', 'autorail']) and not _contains(combined, ['voiture']):
            return 'automotrice', 'diesel'
        if _contains(combined, ['fourgon', 'poste', 'bagages']) and not _contains(combined, ['voiture']):
            return 'wagon', 'none'
        return 'voiture', 'none'

    if _contains(combined, ['wagon', 'wagons', 'citerne', 'citernes', 'silos', 'trémie', 'trémies',
                            'tombereau', 'tombereaux', 'plateau', 'plats', 'couvert', 'couverts',
                            'bâché', 'bache', 'frigorifique', 'minéralier', 'céréalier', 'ciment',
                            'charbon', 'sablière', 'sable', 'ballast', 'traverses', 'affrètement',
                            'affret', 'sncf_f_w', 'france_f_w']):
        return 'wagon', 'none'

    if any(p in path for p in ['sncf_f_le_', '_le_']):
        return 'locomotive', 'electrique'
    if any(p in path for p in ['sncf_f_ld_', '_ld_']):
        return 'locomotive', 'diesel'
    if any(p in path for p in ['sncf_f_lv_', '_lv_']):
        return 'locomotive', 'vapeur'
    if any(p in path for p in ['sncf_f_a_', 'sncf_f_idf_a', 'ratp', 'sncf_f_reg_', 'sncf_f_priv']):
        return 'automotrice', 'electrique'
    if any(p in path for p in ['sncf_f_v_', 'sncf_f_s_']):
        return 'voiture', 'none'
    if any(p in path for p in ['sncf_f_w_', 'france_f_w_', 'affret']):
        return 'wagon', 'none'

    loco_prefixes = ['bb', 'cc', '2d2', '1abba1', '1cc1', '2bb2', '2cc2', 'c ', 'c 20150', 'bbb', 'a1a']
    auto_prefixes = ['x ', 'z ', 'tvg', 'tgv', 'r', 'agc', 'regiolis', 'ter ', 'transilien', 'rer', 'metro']
    coach_prefixes = ['voiture', 'a ', 'b ', 'c ', 'dev', 'uic', 'usi', 'vse', 'corail', 'teoz', 'rib', 'rio', 'rrr', '2n']
    for pr in loco_prefixes:
        if name_l.startswith(pr):
            return 'locomotive', 'electrique'
    for pr in auto_prefixes:
        if name_l.startswith(pr):
            return 'automotrice', 'electrique'
    for pr in coach_prefixes:
        if name_l.startswith(pr):
            return 'voiture', 'none'
    if 'locomotive' in name_l or 'loco' in name_l:
        return 'locomotive', 'electrique'
    if 'automotrice' in name_l or 'autorail' in name_l:
        return 'automotrice', 'diesel'
    return 'wagon', 'none'

# ---------------------------------------------------------------------------
# Entry building
# ---------------------------------------------------------------------------

def compute_price(category, power, passenger_capacity, freight_capacity):
    if category in ('locomotive', 'automotrice'):
        return int(power * 1000) if power else 0
    if category == 'voiture':
        return int(passenger_capacity * 100) if passenger_capacity else 0
    if category == 'wagon':
        return int(freight_capacity * 100) if freight_capacity else 0
    return 0

def status_color(category, max_speed, mass, power, passenger_capacity, freight_capacity):
    if category in ('locomotive', 'automotrice'):
        has_all = max_speed and mass and power
        has_some = max_speed or mass or power
    elif category == 'wagon':
        has_all = max_speed and mass and freight_capacity
        has_some = max_speed or mass or freight_capacity
    elif category == 'voiture':
        has_all = max_speed and mass and passenger_capacity
        has_some = max_speed or mass or passenger_capacity
    else:
        has_all = max_speed and mass
        has_some = max_speed or mass
    if has_all:
        return 'Vert'
    if has_some:
        return 'Jaune/Orange'
    return 'Rouge'

def make_entry(row_id, raw, side, image_path, image_width, specs, is_composite=False):
    nom = raw['nom'] or raw['niv1']
    category, traction = infer_category_traction(raw['url'], raw.get('page_title'), raw['nom'], raw['niv1'], raw['niv2'], raw.get('notes'))
    # TGV / TVG pages are EMU/automotrice
    if 'tgv' in raw['url'].lower() or 'tvg' in nom.lower():
        category = 'automotrice'
    # Steam hardcoded detection
    if 'vapeur' in raw['page_title'].lower() or 'steam' in raw['page_title'].lower():
        traction = 'vapeur'
    length = round(image_width / 10, 1) if image_width else 0.0
    power = int(specs.get('power_kw', 0)) if specs.get('power_kw') else 0
    mass = specs.get('mass_t', 0.0) or 0.0
    max_speed = int(specs.get('max_speed_kmh', 0)) if specs.get('max_speed_kmh') else 0
    first_class = int(specs.get('first_class', 0)) if specs.get('first_class') else 0
    second_class = int(specs.get('second_class', 0)) if specs.get('second_class') else 0
    third_class = int(specs.get('third_class', 0)) if specs.get('third_class') else 0
    passenger_capacity = first_class + second_class + third_class
    freight_capacity = 0
    if category == 'wagon':
        freight_capacity = int(specs.get('tonnage_t', 0)) or 0
    # Tonnage = max payload for wagons, mass for powered, passenger count for coaches
    if category == 'wagon':
        tonnage = freight_capacity
    elif category in ('locomotive', 'automotrice'):
        tonnage = mass
    else:
        tonnage = passenger_capacity
    # price auto
    purchase_price = compute_price(category, power, passenger_capacity, freight_capacity)
    # source string
    sources = []
    if specs.get('_source_te'):
        sources.append('Trains d\'Europe: ' + specs['_source_te'])
    if specs.get('_source_wiki'):
        sources.append(specs['_source_wiki'])
    if specs.get('_source_dbcargo'):
        sources.append('DB Cargo: ' + specs['_source_dbcargo'])
    if not sources:
        sources.append('MLG Traffic (CC BY-NC-SA 3.0)')
    source_str = '; '.join(sources)
    # numberStart
    number_start = ''
    m = re.search(r'\d+', nom)
    if m:
        number_start = m.group(0)
    # seriesName
    series_name = raw['niv2'] or raw['niv1']
    if raw['notes']:
        series_name += ' — ' + raw['notes'] if series_name else raw['notes']
    is_driving = False
    if category == 'voiture' and any(k in (nom + ' ' + raw['notes']).lower() for k in ['pilote','pilot','voiture pilote','b6dux','v2n']):
        is_driving = True
    entry = {
        'id': f'cat-{row_id}',
        'name': nom,
        'category': category,
        'traction': traction,
        'maxSpeed': max_speed,
        'mass': mass,
        'power': power,
        'passengerCapacity': passenger_capacity,
        'freightCapacity': freight_capacity,
        'length': length,
        'imageData': image_path,
        'seriesName': series_name,
        'numberStart': number_start,
        'purchasePrice': purchase_price,
        'cargoTypes': [],
        '_source': source_str,
        'tonnage': tonnage,
    }
    if category == 'wagon':
        xml_name = Path(raw['url']).name
        image_basename = Path(image_path).stem
        sub = guess_wagon_sub(xml_name, image_basename, nom, series_name)
        entry['cargoTypes'] = get_wagon_cargos(sub)
    if is_driving:
        entry['isDrivingTrailer'] = True
    if is_composite:
        entry['_isComposite'] = True
    return entry

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    log('Building Trains d\'Europe indexes...')
    build_te_indexes()

    log('Discovering all MLG collection pages...')
    urls = all_page_urls()
    log(f'Found {len(urls)} MLG pages')

    raw_entries = []
    for url in urls:
        text = fetch(url)
        if not text:
            continue
        if url.lower().endswith('.xml'):
            raw_entries.extend(parse_mlg_xml(url, text))
        else:
            raw_entries.extend(parse_mlg_htm(url, text))
    log(f'Parsed {len(raw_entries)} raw lines')

    # Pre-compute categories and car counts
    raw_meta = []
    for raw in raw_entries:
        if not raw.get('images'):
            continue
        category, traction = infer_category_traction(raw['url'], raw.get('page_title'), raw['nom'], raw['niv1'], raw['niv2'], raw.get('notes'))
        if 'tgv' in raw['url'].lower() or 'tvg' in (raw['nom'] or '').lower():
            category = 'automotrice'
        car_count = len(raw['images']) if category == 'automotrice' else None
        is_composite = len(raw['images']) > 1
        raw_meta.append({'raw': raw, 'category': category, 'traction': traction, 'car_count': car_count, 'is_composite': is_composite})

    # Resolve all individual image bases (R-facing)
    unique_bases = set()
    for m in raw_meta:
        for img in m['raw']['images']:
            unique_bases.add(img)
    unique_bases = list(unique_bases)
    image_cache = {}
    log(f'Resolving {len(unique_bases)} unique images...')
    with ThreadPoolExecutor(max_workers=20) as ex:
        for base in unique_bases:
            image_cache[base] = ex.submit(resolve_image, base)
        done_count = [0]
        for fut in as_completed(image_cache.values()):
            done_count[0] += 1
            if done_count[0] % 500 == 0:
                log(f'  images {done_count[0]}/{len(unique_bases)}')
    for base, fut in image_cache.items():
        try:
            image_cache[base] = fut.result()
        except Exception as e:
            log(f'  image error for {base}: {e}')
            image_cache[base] = []

    # Build composite images
    composite_cache = {}
    log('Building composite images for multi-car entries...')
    for m in raw_meta:
        if not m['is_composite']:
            continue
        raw = m['raw']
        parts = []
        missing = False
        for img in raw['images']:
            res = image_cache.get(img, [])
            if not res:
                missing = True
                break
            parts.append(res[0])
        if missing or not parts:
            continue
        comp_key = tuple(raw['images'])
        if comp_key not in composite_cache:
            comp_path, comp_width = build_composite(raw, parts)
            composite_cache[comp_key] = (comp_path, comp_width)

    # Resolve specs in parallel for unique keys
    spec_cache = {}
    spec_keys = {}
    for m in raw_meta:
        key = (m['category'], m['raw']['nom'], m['raw']['niv1'], m['raw']['niv2'], m['car_count'])
        if key not in spec_keys:
            spec_keys[key] = m
    log(f'Resolving {len(spec_keys)} unique specs...')
    with ThreadPoolExecutor(max_workers=3) as ex:
        future_to_key = {}
        for key, m in spec_keys.items():
            fut = ex.submit(resolve_specs, m['raw']['nom'], m['raw']['niv1'], m['raw']['niv2'], m['category'], m['traction'], m['car_count'], m['raw']['url'])
            future_to_key[fut] = key
        done_count = [0]
        for fut in as_completed(future_to_key):
            done_count[0] += 1
            if done_count[0] % 100 == 0:
                log(f'  specs {done_count[0]}/{len(spec_keys)}')
            key = future_to_key[fut]
            try:
                spec_cache[key] = fut.result() or {}
            except Exception as e:
                log(f'  spec error {key}: {e}')
                spec_cache[key] = {}

    # Process entries
    entries = []
    excel_rows = []
    for i, m in enumerate(raw_meta, 1):
        raw = m['raw']
        country, _ = country_from_url(raw['url'])
        if m['is_composite']:
            comp_key = tuple(raw['images'])
            comp_path, comp_width = composite_cache.get(comp_key, (None, 0))
            if not comp_path:
                log(f'  no composite for {raw["nom"]}')
                continue
            side = ''
            img_path = comp_path
            width = comp_width
            is_composite = True
        else:
            res = image_cache.get(raw['images'][0], [])
            if not res:
                log(f'  no image for {raw["nom"]} ({raw["images"][0]})')
                continue
            side, img_path, width = res[0]
            is_composite = False
        key = (m['category'], raw['nom'], raw['niv1'], raw['niv2'], m['car_count'])
        specs = spec_cache.get(key, {})
        row_id = len(entries) + 1
        entry = make_entry(row_id, raw, side, img_path, width, specs, is_composite=is_composite)
        entries.append(entry)
        excel_rows.append({
            'ID': entry['id'],
            'Pays': country,
            'Page MLG': raw['url'].split('/')[-1],
            'Groupe': raw['niv1'],
            'Sous-groupe': raw['niv2'],
            'Nom': entry['name'],
            'Image': Path(entry['imageData']).name,
            'Côté': side,
            'Catégorie': entry['category'],
            'Traction': entry['traction'],
            'Vitesse': entry['maxSpeed'],
            'Masse': entry['mass'],
            'Puissance': entry['power'],
            'Longueur': entry['length'],
            'Places': entry['passengerCapacity'],
            'Tonnage': entry['freightCapacity'],
            'Prix': entry['purchasePrice'],
            'Source': entry['_source'],
            'Statut': status_color(entry['category'], entry['maxSpeed'], entry['mass'], entry['power'], entry['passengerCapacity'], entry['freightCapacity']),
            'Notes': raw['notes'],
        })
        if i % 500 == 0:
            log(f'  processed {i}/{len(raw_meta)} -> {len(entries)} catalog entries')

    log(f'Final catalog entries: {len(entries)}')

    # Write catalog-data.js (keep cargo types header, replace CATALOG)
    header = load_catalog_header()
    catalog_json = json.dumps(entries, indent=2, ensure_ascii=False)
    final = header + '\n\nexport const CATALOG = ' + catalog_json + ';\n'
    write_text(CATALOG_PATH, final)
    log(f'Wrote {CATALOG_PATH}')

    # Write Excel
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Catalogue MLG'
    if excel_rows:
        headers = list(excel_rows[0].keys())
        ws.append(headers)
        color_map = {
            'Vert': 'C6EFCE',
            'Jaune/Orange': 'FFEB9C',
            'Rouge': 'FFC7CE',
        }
        for r in excel_rows:
            row = [r[h] for h in headers]
            ws.append(row)
            fill = PatternFill(start_color=color_map.get(r['Statut'], 'FFFFFF'), end_color=color_map.get(r['Statut'], 'FFFFFF'), fill_type='solid')
            ws.cell(row=ws.max_row, column=headers.index('Statut') + 1).fill = fill
    else:
        ws.append(['Aucune entrée'])
    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(EXCEL_PATH)
    log(f'Wrote {EXCEL_PATH}')

if __name__ == '__main__':
    main()
