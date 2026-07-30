#!/usr/bin/env python3
# Reconstruction complète du catalogue France à partir de MLG Traffic + recherche web
# Usage: python3 scripts/rebuild_france.py
# Génère js/catalog-data.js (partie CATALOG) et catalog_france.xlsx

import json, re, os, sys, math, time, urllib.request, urllib.error, urllib.parse, html, threading
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed
from PIL import Image
from io import BytesIO
from bs4 import BeautifulSoup

REPO = Path('/home/ubuntu/repos/RAIL-EMPIRE-APP')
IMG_ROOT = REPO / 'img/catalog'
CATALOG_PATH = REPO / 'js/catalog-data.js'
EXCEL_PATH = REPO / 'catalog_france.xlsx'
CACHE_DIR = REPO / '.cache_rebuild'
MLG_BASE = 'http://www.mlgtraffic.net'

COUNTRY = 'France'
COLLECTION_URL = f'{MLG_BASE}/Coll_FR_F.htm'

CACHE_DIR.mkdir(exist_ok=True)
fetch_lock = threading.Lock()


def log(msg):
    print(msg, flush=True)

def slugify(s):
    s = re.sub(r'[^\w\s]', ' ', s)
    return re.sub(r'\s+', ' ', s).strip().lower()

def norm_digits(s):
    return set(re.findall(r'\d+', str(s)))

def read_text(path):
    with open(path, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read()

def write_text(path, text):
    with open(path, 'w', encoding='utf-8') as f:
        f.write(text)

def fetch(url, timeout=30):
    """Fetch with caching (thread-safe)."""
    safe = re.sub(r'[^\w\-_.]', '_', url)
    cache = CACHE_DIR / safe
    with fetch_lock:
        if cache.exists() and cache.stat().st_size > 0:
            return read_text(cache)
    try:
        req = urllib.request.Request(url, headers={'User-Agent':'DevinBot/1.0'})
        data = urllib.request.urlopen(req, timeout=timeout).read()
        text = data.decode('utf-8', errors='ignore')
        with fetch_lock:
            # write only if still missing (another thread may have written)
            if not cache.exists():
                write_text(cache, text)
        return text
    except Exception as e:
        log(f'  fetch fail {url}: {e}')
        return None

def fetch_bytes(url, timeout=20):
    try:
        req = urllib.request.Request(url, headers={'User-Agent':'DevinBot/1.0'})
        return urllib.request.urlopen(req, timeout=timeout).read()
    except Exception as e:
        return None

def load_catalog_header():
    """Load CATALOG_CARGO_TYPES header from existing file or default."""
    if CATALOG_PATH.exists():
        text = read_text(CATALOG_PATH)
        m = re.search(r'export const CATALOG_CARGO_TYPES = \[.*?\];', text, re.S)
        if m:
            return m.group(0)
    # fallback minimal cargo types
    return '''// AUTO-GENERATED — Catalogue complet MLG Traffic
export const CATALOG_CARGO_TYPES = [
  {"category":"marchandises","type":"general","name":"Marchandises générales","unit":"t","pricePerUnit":20,"hazard":false}
];
'''

# ---------------------------------------------------------------------------
# Trains d'Europe index & spec resolver
# ---------------------------------------------------------------------------

te_indexes = {}

def build_te_indexes():
    """Fetch all Trains d'Europe indexes and build a slug map."""
    indexes = {
        'locomotives': 'http://trains-europe.fr/sncf/locomotives/index.htm',
        'automoteurs': 'http://trains-europe.fr/sncf/automoteurs/index.htm',
        'tgv': 'http://trains-europe.fr/sncf/tgv/index.htm',
        'turbo': 'http://trains-europe.fr/sncf/turbo/index.htm',
        'voitures': 'http://trains-europe.fr/sncf/voitures/index.htm',
    }
    for name, url in indexes.items():
        text = fetch(url)
        if not text:
            continue
        soup = BeautifulSoup(text, 'html.parser')
        entries = []
        for a in soup.find_all('a', href=True):
            href = a['href'].strip()
            txt = a.get_text(' ', strip=True)
            if not (href.endswith('.htm') or href.endswith('.html')):
                continue
            if href.startswith('..') or href.startswith('http'):
                continue
            if 'index' in href.lower():
                continue
            # resolve relative to base url
            base = url.rsplit('/', 1)[0]
            full = f'{base}/{href}'
            entries.append({'title': txt, 'slug': full, 'href': href})
        te_indexes[name] = entries
        log(f'TE index {name}: {len(entries)} pages')

def te_find_slug(name, category):
    """Find best matching Trains d'Europe slug for a given name/category."""
    if not name:
        return None
    norm = slugify(name)
    tokens = set(norm.split())
    digits = norm_digits(name)
    candidates = []
    if category in ('locomotive', 'automotrice', 'tgv'):
        for idx_name in (('tgv' if 'tgv' in norm else None),
                         ('turbo' if 'turbo' in norm or 'rtg' in norm or 'etg' in norm else None),
                         ('automoteurs' if category == 'automotrice' else None),
                         ('locomotives' if category in ('locomotive','automotrice') else None)):
            if not idx_name:
                continue
            for e in te_indexes.get(idx_name, []):
                title_norm = slugify(e['title'])
                title_digits = norm_digits(e['title'])
                score = 0
                # exact class token match
                if norm and (norm in title_norm or title_norm in norm):
                    score += 50
                # digit overlap
                if digits and title_digits and digits & title_digits:
                    score += 20
                # keyword matches
                for tok in tokens:
                    if tok in title_norm:
                        score += 5
                # prefer TGV index for TGV
                if idx_name == 'tgv' and 'tgv' in norm:
                    score += 30
                if score:
                    candidates.append((score, e))
    else:
        for e in te_indexes.get('voitures', []):
            title_norm = slugify(e['title'])
            title_digits = norm_digits(e['title'])
            score = 0
            if norm in title_norm or title_norm in norm:
                score += 50
            if digits and title_digits and digits & title_digits:
                score += 20
            for tok in tokens:
                if tok in title_norm:
                    score += 5
            if score:
                candidates.append((score, e))
    if not candidates:
        return None
    candidates.sort(key=lambda x: x[0], reverse=True)
    return candidates[0][1]['slug']

te_spec_cache = {}
te_cache_lock = threading.Lock()

def parse_te_page(url):
    """Parse a Trains d'Europe page into a list of fiche dicts."""
    with te_cache_lock:
        if url in te_spec_cache:
            return te_spec_cache[url]
    text = fetch(url)
    if not text:
        return []
    soup = BeautifulSoup(text, 'html.parser')
    fiches = []
    for fiche in soup.find_all(['div', 'td'], class_=lambda c: c and any('fiche' in cls for cls in c)):
        data = {}
        for d in fiche.find_all('div', style=lambda s: s and 'display:inline' in s):
            strong = d.find('div', style=lambda s: s and 'font-weight:bold' in s)
            if not strong:
                continue
            label = strong.get_text(strip=True).rstrip(':').rstrip()
            value = d.get_text(' ', strip=True).replace(strong.get_text(strip=True), '', 1).strip()
            data[label] = value
        if data:
            fiches.append(data)
    with te_cache_lock:
        te_spec_cache[url] = fiches
    return fiches

def norm_num_str(s):
    """Collapse all visual spacing characters between digits."""
    s2 = str(s)
    # collapse unicode non-breaking spaces / thin spaces between digits
    for ch in '\u00a0\u2007\u202f\u2009\u2006\u2002\u2003\u2008\u205f':
        s2 = s2.replace(ch, ' ')
    s2 = re.sub(r'(?<=\d)\s+(?=\d)', '', s2)
    return s2

def first_num(s):
    if not s:
        return None
    s2 = norm_num_str(s)
    m = re.search(r'[\d\s]+(?:[,.][\d\s]+)?', s2)
    if not m:
        return None
    v = m.group(0).replace(' ', '').replace(',', '.')
    try:
        return float(v)
    except Exception:
        return None

def all_nums(s):
    if not s:
        return []
    s2 = norm_num_str(s)
    vals = []
    for m in re.finditer(r'\d[\d\s,\.]*\d|\d', s2):
        v = m.group(0).replace(' ', '').replace(',', '.')
        try:
            vals.append(float(v))
        except Exception:
            pass
    return vals

def pick_fiche(fiches, name, niv2):
    """Pick the best fiche matching a subseries."""
    if not fiches:
        return None
    if len(fiches) == 1:
        return fiches[0]
    # try to match by series indicators in name/niv2
    series_indicators = []
    for m in re.finditer(r'(1[èe]re?|premi[èe]re|2[èe]me|deuxi[èe]me|3[èe]me|troisi[èe]me|4[èe]me)', str(name)+' '+str(niv2), re.I):
        series_indicators.append(m.group(1).lower()[:3])
    if not series_indicators:
        return fiches[0]
    # fiche contains Numéros or text
    for i, f in enumerate(fiches):
        nums = str(f.get('Numéros', '')) + ' ' + str(f.get('Nombre d\'exemplaires', ''))
        if series_indicators[0] in nums.lower():
            return f
    return fiches[0]

def parse_spec_value(label, value):
    """Extract numeric / text fields from a fiche."""
    if not value:
        return {}
    res = {}
    vstr = str(value)
    if label in ('Puissance',):
        nums = all_nums(vstr)
        if nums:
            res['power_kw'] = max(nums) if 'sous' in vstr.lower() or any(x in vstr for x in ['ou','et',';']) else nums[0]
    elif label in ('Poids', 'Masse', 'Masse en service', 'Masse totale'):
        nums = all_nums(vstr)
        if nums:
            res['mass_t'] = nums[0]
    elif label in ('Vitesse maximale', 'Vitesse'):
        nums = all_nums(vstr)
        if nums:
            res['max_speed_kmh'] = max(nums)
    elif label in ('Longueur', 'Longueur hors-tout', 'Longueur totale'):
        nums = all_nums(vstr)
        if nums:
            res['length_m'] = nums[0]
    elif label in ('Places de première classe',):
        nums = all_nums(vstr)
        if nums:
            res['first_class'] = int(nums[0])
    elif label in ('Places de seconde classe',):
        nums = all_nums(vstr)
        if nums:
            res['second_class'] = int(nums[0])
    elif label in ('Places', 'Capacité'):
        nums = all_nums(vstr)
        if nums:
            res['capacity'] = int(nums[0])
    return res

def merge_specs(fiche, name, niv2):
    spec = {}
    if fiche:
        for label, val in fiche.items():
            spec.update(parse_spec_value(label, val))
    return spec

def resolve_specs_te(name, niv1, niv2, category, traction):
    """Try to resolve specs from Trains d'Europe."""
    # Use Niv1 title for TGV, otherwise name
    search_name = name if category != 'tgv' else (niv1 or name)
    slug = te_find_slug(search_name, category)
    if not slug:
        return {}
    fiches = parse_te_page(slug)
    if not fiches:
        return {}
    f = pick_fiche(fiches, name, niv2)
    spec = merge_specs(f, name, niv2)
    if not spec:
        return {}
    spec['_source_te'] = slug
    return spec

# ---------------------------------------------------------------------------
# Wikipedia FR fallback (for old classes, non-wagons)
# ---------------------------------------------------------------------------

def resolve_specs_wikipedia(name, category):
    """Fetch French Wikipedia page for the class and parse infobox/text."""
    if category == 'wagon':
        return {}  # user forbids Wikipedia for wagons
    if not name:
        return {}
    # Build wiki title: e.g. BB 7200, 2D2 5500
    wiki = name.split()[0] + '_' + name.split()[1] if len(name.split()) >= 2 else name.replace(' ', '_')
    url = f'https://fr.wikipedia.org/wiki/{wiki}'
    cache_key = re.sub(r'[^\w]', '_', wiki)
    cache = CACHE_DIR / f'wiki_{cache_key}'
    if cache.exists() and cache.stat().st_size > 0:
        text = read_text(cache)
    else:
        data = fetch_bytes(url)
        if not data:
            return {}
        text = data.decode('utf-8', errors='ignore')
        write_text(cache, text)
    soup = BeautifulSoup(text, 'html.parser')
    spec = {}
    # Try infobox table
    for table in soup.find_all('table'):
        rows = []
        for tr in table.find_all('tr'):
            cells = [c.get_text(' ', strip=True) for c in tr.find_all(['th','td'])]
            if len(cells) >= 2:
                rows.append(cells)
        for cells in rows:
            k = cells[0]
            v = ' '.join(cells[1:])
            kl = k.lower()
            if 'puissance' in kl:
                nums = all_nums(v)
                if nums and 'power_kw' not in spec:
                    spec['power_kw'] = max(nums)
            elif 'masse' in kl and 'en service' in kl:
                nums = all_nums(v)
                if nums and 'mass_t' not in spec:
                    spec['mass_t'] = nums[0]
            elif 'vitesse maximale' in kl:
                nums = all_nums(v)
                if nums:
                    spec['max_speed_kmh'] = max(nums)
            elif 'longueur' in kl:
                nums = all_nums(v)
                if nums and 'length_m' not in spec:
                    spec['length_m'] = nums[0]
    # Fallback text regex
    body = soup.get_text(' ', strip=True)
    if 'power_kw' not in spec:
        m = re.search(r'(?i)puissance[^.\n]{0,80}?(\d[\d\s,\.]+)\s*kW', body)
        if m:
            nums = all_nums(m.group(1))
            if nums:
                spec['power_kw'] = max(nums)
    if 'mass_t' not in spec:
        m = re.search(r'(?i)masse en service[^.\n]{0,80}?(\d[\d\s,\.]+)\s*t', body)
        if m:
            nums = all_nums(m.group(1))
            if nums:
                spec['mass_t'] = nums[0]
    if 'max_speed_kmh' not in spec:
        m = re.search(r'(?i)vitesse maximale[^.\n]{0,80}?(\d[\d\s,\.]+)\s*km/h', body)
        if m:
            nums = all_nums(m.group(1))
            if nums:
                spec['max_speed_kmh'] = max(nums)
    if spec:
        spec['_source_wiki'] = url
    return spec

# ---------------------------------------------------------------------------
# MLG parsing
# ---------------------------------------------------------------------------

def france_page_urls():
    """Return list of France MLG page URLs from collection index."""
    text = fetch(COLLECTION_URL)
    if not text:
        return []
    urls = []
    for m in re.finditer(r'href="([^"]+)"', text):
        href = m.group(1)
        # remove anchor
        href = href.split('#')[0]
        # determine extension ignoring query
        base_href = href.split('?')[0]
        if base_href.endswith('.xml'):
            href = base_href
        if base_href.endswith('.xml') or base_href.endswith('.htm'):
            # encode spaces and special chars in path while preserving query string
            if '?' in href:
                path, qs = href.split('?', 1)
                path = urllib.parse.quote(path, safe='/')
                href = f'{path}?{qs}'
            else:
                href = urllib.parse.quote(href, safe='/')
            urls.append(f'{MLG_BASE}/{href}')
    return urls

def parse_mlg_xml(url, text):
    """Parse an MLG XML file into raw entries."""
    import xml.etree.ElementTree as ET
    root = ET.fromstring(text)
    page_title = ''
    for tag in ['Nompage1', 'Nompage2']:
        el = root.find(tag)
        if el is not None and el.text:
            page_title += ' ' + el.text.strip()
    entries = []
    for niv1 in root.findall('Niv1'):
        t1 = niv1.find('Titre1')
        niv1_title = (t1.text or '').strip() if t1 is not None else ''
        for niv2 in niv1.findall('Niv2'):
            t2 = niv2.find('Titre2')
            niv2_title = (t2.text or '').strip() if t2 is not None else ''
            for ligne in niv2.findall('Ligne'):
                nom_el = ligne.find('Nom')
                img_el = ligne.find('Image')
                notes_el = ligne.find('Notes')
                nom = (nom_el.text or '').strip() if nom_el is not None and nom_el.text else ''
                image = (img_el.text or '').strip() if img_el is not None and img_el.text else ''
                notes = (notes_el.text or '').strip() if notes_el is not None and notes_el.text else ''
                if not image:
                    continue
                entries.append({
                    'url': url,
                    'page_title': page_title.strip(),
                    'niv1': niv1_title,
                    'niv2': niv2_title,
                    'nom': nom,
                    'image': image,
                    'notes': notes,
                })
    return entries

def parse_mlg_htm(url, text):
    """Parse an MLG HTML data page (const dataPage JSON) into raw entries."""
    text = text.replace('\r', '\n')
    m = re.search(r'const\s+dataPage\s*=\s*(\{.*?\}\s*);?\s*<\/script>', text, re.S)
    if not m:
        return []
    raw = m.group(1)
    # remove trailing commas
    raw = re.sub(r',(\s*[}\]])', r'\1', raw)
    try:
        data = json.loads(raw)
    except Exception as e:
        log(f'JSON parse fail {url}: {e}')
        return []
    page_title = ' '.join([str(data.get(k, '')) for k in ['titrePageF','titre2F','titre3F'] if data.get(k)])
    entries = []
    niv1_list = data.get('niv1', [])
    if isinstance(niv1_list, dict):
        niv1_list = [niv1_list]
    for niv1 in niv1_list:
        niv1_title = str(niv1.get('libF', ''))
        niv2_list = niv1.get('niv2', [])
        if isinstance(niv2_list, dict):
            niv2_list = [niv2_list]
        for niv2 in (niv2_list or [niv1]):
            niv2_title = str(niv2.get('libF', '') or niv2.get('txtF', ''))
            mat_list = niv2.get('mat', [])
            if isinstance(mat_list, dict):
                mat_list = [mat_list]
            for mat in mat_list:
                image = str(mat.get('image', ''))
                nom = str(mat.get('nomF', '') or mat.get('txtF', '') or niv1_title)
                notes = str(mat.get('descrF', ''))
                if not image:
                    continue
                entries.append({
                    'url': url,
                    'page_title': page_title.strip(),
                    'niv1': niv1_title,
                    'niv2': niv2_title,
                    'nom': nom,
                    'image': image,
                    'notes': notes,
                })
    return entries

# ---------------------------------------------------------------------------
# Image handling
# ---------------------------------------------------------------------------

def resolve_image(base_rel):
    """Download the right-facing image variant and return a single (side, local_relative_path, width)."""
    base_name = Path(base_rel).name
    sub_dir = Path(base_rel).parent.relative_to('images') if base_rel.startswith('images/') else Path('.')
    local_dir = IMG_ROOT / sub_dir
    local_dir.mkdir(parents=True, exist_ok=True)
    # Priority: right-facing variants, then non-sided, then left-facing as last resort.
    priority = [
        ('R', f'{base_name}_R.gif'),
        ('R', f'{base_name}_Anim_R.gif'),
        ('', f'{base_name}.gif'),
        ('', f'{base_name}_Anim.gif'),
        ('L', f'{base_name}_L.gif'),
        ('L', f'{base_name}_Anim_L.gif'),
    ]
    encoded_base = urllib.parse.quote(base_rel, safe='/')
    for side, filename in priority:
        local = local_dir / filename
        if local.exists() and local.stat().st_size > 0:
            try:
                im = Image.open(local)
                return [(side or 'R', str(local.relative_to(REPO)), im.width)]
            except Exception:
                continue
        suffix = filename[len(base_name):] if filename.startswith(base_name) else filename
        url = f'{MLG_BASE}/{encoded_base}{urllib.parse.quote(suffix, safe="/")}'
        data = fetch_bytes(url, timeout=8)
        if data:
            local.write_bytes(data)
            try:
                im = Image.open(BytesIO(data))
                return [(side or 'R', str(local.relative_to(REPO)), im.width)]
            except Exception:
                continue
    return []

# ---------------------------------------------------------------------------
# Categorisation
# ---------------------------------------------------------------------------

def infer_category_traction(page_url, nom, niv1, niv2):
    path = page_url.split('/')[-1].lower()
    name_l = nom.lower()
    niv1_l = niv1.lower()
    niv2_l = niv2.lower()
    if 'le_' in path or 'locomotive' in niv1_l + niv2_l or path.startswith('sncf_f_le_'):
        return 'locomotive', 'electrique'
    if 'ld_' in path or 'diesel' in niv1_l + niv2_l or path.startswith('sncf_f_ld_'):
        return 'locomotive', 'diesel'
    if 'lv_' in path or 'vapeur' in niv1_l + niv2_l or 'steam' in niv1_l:
        return 'locomotive', 'vapeur'
    if any(p in path for p in ['ae_','ael','a_ter','ad_','turbo','tgv','xr','sncf_f_idf_a','ratp','sncf_f_reg_','sncf_f_priv','france_f_trv','hexafret']):
        return 'automotrice', 'electrique'
    if any(p in path for p in ['sncf_f_v_','france_f_poste','sncf_f_s_']):
        if 'draisine' in name_l or 'automotrice' in name_l:
            return 'automotrice', 'diesel'
        if 'fourgon' in name_l or 'poste' in name_l or 's_' in path:
            return 'wagon', 'none'
        return 'voiture', 'none'
    if any(p in path for p in ['sncf_f_w','france_f_w','affret']):
        return 'wagon', 'none'
    # fallback by name prefix
    loco_prefixes = ['bb','cc','2d2','1abba1','1cc1','2bb2','2cc2','c ','c 20150','bbb','a1a']
    auto_prefixes = ['x ','z ','tvg','tgv','r','agc','regiolis','ter ','transilien','rer','metro']
    coach_prefixes = ['voiture','a ','b ','c ','dev','uic','usi','vse','corail','teoz','rib','rio','rrr','2n']
    for pr in loco_prefixes:
        if name_l.startswith(pr):
            return 'locomotive', 'electrique'
    for pr in auto_prefixes:
        if name_l.startswith(pr):
            return 'automotrice', 'electrique'
    for pr in coach_prefixes:
        if name_l.startswith(pr):
            return 'voiture', 'none'
    if 'locomotive' in name_l or 'loco' in name_l:
        return 'locomotive', 'electrique'
    if 'automotrice' in name_l or 'autorail' in name_l:
        return 'automotrice', 'diesel'
    return 'wagon', 'none'

# ---------------------------------------------------------------------------
# Entry building
# ---------------------------------------------------------------------------

def compute_price(category, power, passenger_capacity, freight_capacity):
    if category in ('locomotive', 'automotrice'):
        return int(power * 1000) if power else 0
    if category == 'voiture':
        return int(passenger_capacity * 100) if passenger_capacity else 0
    if category == 'wagon':
        return int(freight_capacity * 100) if freight_capacity else 0
    return 0

def status_color(category, max_speed, mass, power, passenger_capacity, freight_capacity):
    if category in ('locomotive', 'automotrice'):
        has_all = max_speed and mass and power
        has_some = max_speed or mass or power
    elif category == 'wagon':
        has_all = max_speed and mass and freight_capacity
        has_some = max_speed or mass or freight_capacity
    elif category == 'voiture':
        has_all = max_speed and mass and passenger_capacity
        has_some = max_speed or mass or passenger_capacity
    else:
        has_all = max_speed and mass
        has_some = max_speed or mass
    if has_all:
        return 'Vert'
    if has_some:
        return 'Jaune/Orange'
    return 'Rouge'

def make_entry(row_id, raw, side, image_path, image_width, specs):
    nom = raw['nom'] or raw['niv1']
    category, traction = infer_category_traction(raw['url'], raw['nom'], raw['niv1'], raw['niv2'])
    # Special cases: TGV page Ligne names are car types, category should be automotrice for train
    if 'sncf_f_tgv' in raw['url'].lower() or 'tvg' in nom.lower():
        category = 'automotrice'
    # Steam hardcoded detection
    if 'vapeur' in raw['page_title'].lower() or 'steam' in raw['page_title'].lower():
        traction = 'vapeur'
    length = round(image_width / 10, 1) if image_width else 0.0
    power = int(specs.get('power_kw', 0)) if specs.get('power_kw') else 0
    mass = specs.get('mass_t', 0.0) or 0.0
    max_speed = int(specs.get('max_speed_kmh', 0)) if specs.get('max_speed_kmh') else 0
    first_class = int(specs.get('first_class', 0)) if specs.get('first_class') else 0
    second_class = int(specs.get('second_class', 0)) if specs.get('second_class') else 0
    passenger_capacity = first_class + second_class
    freight_capacity = 0
    if category == 'wagon':
        freight_capacity = int(specs.get('tonnage_t', 0)) or 0
    tonnage = freight_capacity if category == 'wagon' else (mass if category in ('locomotive','automotrice') else passenger_capacity)
    # price auto
    purchase_price = compute_price(category, power, passenger_capacity, freight_capacity)
    # source string
    sources = []
    if specs.get('_source_te'):
        sources.append('Trains d\'Europe: ' + specs['_source_te'])
    if specs.get('_source_wiki'):
        sources.append(specs['_source_wiki'])
    if not sources:
        sources.append('MLG Traffic (CC BY-NC-SA 3.0)')
    source_str = '; '.join(sources)
    # numberStart
    number_start = ''
    m = re.search(r'\d+', nom)
    if m:
        number_start = m.group(0)
    # seriesName
    series_name = raw['niv2'] or raw['niv1']
    if raw['notes']:
        series_name += ' — ' + raw['notes'] if series_name else raw['notes']
    is_driving = False
    if category == 'voiture' and any(k in (nom + ' ' + raw['notes']).lower() for k in ['pilote','pilot','voiture pilote','b6dux','v2n']):
        is_driving = True
    entry = {
        'id': f'cat-{row_id}',
        'name': nom,
        'category': category,
        'traction': traction,
        'maxSpeed': max_speed,
        'mass': mass,
        'power': power,
        'passengerCapacity': passenger_capacity,
        'freightCapacity': freight_capacity,
        'length': length,
        'imageData': image_path,
        'seriesName': series_name,
        'numberStart': number_start,
        'purchasePrice': purchase_price,
        'cargoTypes': [],
        '_source': source_str,
        'tonnage': tonnage,
    }
    if is_driving:
        entry['isDrivingTrailer'] = True
    return entry

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    log('Building Trains d\'Europe indexes...')
    build_te_indexes()

    log('Fetching France collection pages...')
    urls = france_page_urls()
    log(f'Found {len(urls)} France pages')

    raw_entries = []
    for url in urls:
        text = fetch(url)
        if not text:
            continue
        if url.endswith('.xml'):
            raw_entries.extend(parse_mlg_xml(url, text))
        else:
            raw_entries.extend(parse_mlg_htm(url, text))
    log(f'Parsed {len(raw_entries)} raw lines')

    # Pre-compute categories and unique spec keys
    raw_meta = []
    for raw in raw_entries:
        if not raw['image']:
            continue
        category, traction = infer_category_traction(raw['url'], raw['nom'], raw['niv1'], raw['niv2'])
        if 'sncf_f_tgv' in raw['url'].lower():
            category = 'automotrice'
        raw_meta.append({'raw': raw, 'category': category, 'traction': traction})

    # Resolve images in parallel for unique image bases
    unique_bases = list(set(r['raw']['image'] for r in raw_meta))
    image_cache = {}
    log(f'Resolving {len(unique_bases)} unique images...')
    with ThreadPoolExecutor(max_workers=20) as ex:
        for i, base in enumerate(unique_bases):
            future = ex.submit(resolve_image, base)
            image_cache[base] = future
        # consume as completed with progress
        done_count = [0]
        for fut in as_completed(image_cache.values()):
            done_count[0] += 1
            if done_count[0] % 100 == 0:
                log(f'  images {done_count[0]}/{len(unique_bases)}')
    # unwrap futures
    for base, fut in image_cache.items():
        try:
            image_cache[base] = fut.result()
        except Exception as e:
            log(f'  image error for {base}: {e}')
            image_cache[base] = []

    # Resolve specs in parallel for unique (category,name,niv1,niv2) keys
    spec_cache = {}
    spec_keys = {}
    for m in raw_meta:
        key = (m['category'], m['raw']['nom'], m['raw']['niv1'], m['raw']['niv2'])
        if key not in spec_keys:
            spec_keys[key] = m
    log(f'Resolving {len(spec_keys)} unique specs...')
    with ThreadPoolExecutor(max_workers=10) as ex:
        future_to_key = {}
        for key, m in spec_keys.items():
            fut = ex.submit(resolve_specs_te, m['raw']['nom'], m['raw']['niv1'], m['raw']['niv2'], m['category'], m['traction'])
            future_to_key[fut] = key
        for fut in as_completed(future_to_key):
            key = future_to_key[fut]
            try:
                spec = fut.result()
                if not spec:
                    # try Wikipedia fallback (non-wagons only) in same thread
                    m = spec_keys[key]
                    if m['category'] != 'wagon':
                        spec = resolve_specs_wikipedia(m['raw']['nom'], m['category'])
                spec_cache[key] = spec or {}
            except Exception as e:
                log(f'  spec error {key}: {e}')
                spec_cache[key] = {}

    # Process entries
    entries = []
    excel_rows = []
    for i, m in enumerate(raw_meta, 1):
        raw = m['raw']
        images = image_cache.get(raw['image'], [])
        if not images:
            log(f'  no image for {raw["nom"]} ({raw["image"]})')
            continue
        key = (m['category'], raw['nom'], raw['niv1'], raw['niv2'])
        specs = spec_cache.get(key, {})
        for side, img_path, width in images:
            row_id = len(entries) + 1
            entry = make_entry(row_id, raw, side, img_path, width, specs)
            entries.append(entry)
            excel_rows.append({
                'ID': entry['id'],
                'Pays': COUNTRY,
                'Page MLG': raw['url'].split('/')[-1],
                'Groupe': raw['niv1'],
                'Sous-groupe': raw['niv2'],
                'Nom': entry['name'],
                'Image': Path(entry['imageData']).name,
                'Côté': side,
                'Catégorie': entry['category'],
                'Traction': entry['traction'],
                'Vitesse': entry['maxSpeed'],
                'Masse': entry['mass'],
                'Puissance': entry['power'],
                'Longueur': entry['length'],
                'Places': entry['passengerCapacity'],
                'Tonnage': entry['freightCapacity'],
                'Prix': entry['purchasePrice'],
                'Source': entry['_source'],
                'Statut': status_color(entry['category'], entry['maxSpeed'], entry['mass'], entry['power'], entry['passengerCapacity'], entry['freightCapacity']),
                'Notes': raw['notes'],
            })
        if i % 100 == 0:
            log(f'  processed {i}/{len(raw_meta)} -> {len(entries)} catalog entries')

    log(f'Final catalog entries: {len(entries)}')

    # Write catalog-data.js (keep cargo types header, replace CATALOG)
    header = load_catalog_header()
    catalog_json = json.dumps(entries, indent=2, ensure_ascii=False)
    final = header + '\n\nexport const CATALOG = ' + catalog_json + ';\n'
    write_text(CATALOG_PATH, final)
    log(f'Wrote {CATALOG_PATH}')

    # Write Excel
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Catalogue France'
    if excel_rows:
        headers = list(excel_rows[0].keys())
        ws.append(headers)
        color_map = {
            'Vert': 'C6EFCE',
            'Jaune/Orange': 'FFEB9C',
            'Rouge': 'FFC7CE',
        }
        for r in excel_rows:
            row = [r[h] for h in headers]
            ws.append(row)
            fill = PatternFill(start_color=color_map.get(r['Statut'], 'FFFFFF'), end_color=color_map.get(r['Statut'], 'FFFFFF'), fill_type='solid')
            ws.cell(row=ws.max_row, column=headers.index('Statut') + 1).fill = fill
    else:
        ws.append(['Aucune entrée'])
    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(EXCEL_PATH)
    log(f'Wrote {EXCEL_PATH}')

if __name__ == '__main__':
    main()
